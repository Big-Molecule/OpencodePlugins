#!/usr/bin/env python3
"""Append a paragraph (optional heading level) to a DOCX."""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path


def main() -> int:
    p = argparse.ArgumentParser()
    p.add_argument("docx")
    p.add_argument("--text", required=True)
    p.add_argument("--heading-level", type=int, default=0, help="0=body, 1-3=heading")
    p.add_argument("--out", default="")
    args = p.parse_args()

    path = Path(args.docx)
    if not path.is_file():
        print(json.dumps({"ok": False, "error": f"not found: {path}"}))
        return 1

    from docx import Document

    doc = Document(str(path))
    if args.heading_level and 1 <= args.heading_level <= 3:
        doc.add_heading(args.text, level=args.heading_level)
    else:
        doc.add_paragraph(args.text)

    out = Path(args.out) if args.out.strip() else path
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    print(json.dumps({"ok": True, "path": str(out.resolve())}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
