#!/usr/bin/env python3
"""Set header/footer text for all sections."""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path


def main() -> int:
    p = argparse.ArgumentParser()
    p.add_argument("docx")
    p.add_argument("--header", default="")
    p.add_argument("--footer", default="")
    p.add_argument("--out", default="")
    args = p.parse_args()

    path = Path(args.docx)
    if not path.is_file():
        print(json.dumps({"ok": False, "error": f"not found: {path}"}))
        return 1

    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document(str(path))
    for section in doc.sections:
        if args.header != "":
            header = section.header
            header.is_linked_to_previous = False
            if header.paragraphs:
                header.paragraphs[0].text = args.header
            else:
                header.add_paragraph(args.header)
        if args.footer != "":
            footer = section.footer
            footer.is_linked_to_previous = False
            if footer.paragraphs:
                p0 = footer.paragraphs[0]
                p0.text = args.footer
                p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                fp = footer.add_paragraph(args.footer)
                fp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    out = Path(args.out) if args.out.strip() else path
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    print(json.dumps({"ok": True, "path": str(out.resolve())}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
