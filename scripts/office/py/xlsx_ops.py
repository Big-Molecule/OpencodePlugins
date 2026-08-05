#!/usr/bin/env python3
"""Basic XLSX read/write via openpyxl."""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path


def main() -> int:
    p = argparse.ArgumentParser()
    p.add_argument("xlsx")
    p.add_argument("--action", choices=["info", "read", "write_cell", "create"], required=True)
    p.add_argument("--sheet", default="")
    p.add_argument("--cell", default="A1")
    p.add_argument("--value", default="")
    p.add_argument("--max-rows", type=int, default=50)
    p.add_argument("--max-cols", type=int, default=20)
    p.add_argument("--out", default="")
    args = p.parse_args()

    from openpyxl import Workbook, load_workbook

    path = Path(args.xlsx)

    if args.action == "create":
        wb = Workbook()
        ws = wb.active
        ws.title = args.sheet or "Sheet1"
        out = Path(args.out) if args.out.strip() else path
        out.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(out))
        print(json.dumps({"ok": True, "path": str(out.resolve())}, ensure_ascii=False))
        return 0

    if not path.is_file():
        print(json.dumps({"ok": False, "error": f"not found: {path}"}))
        return 1

    wb = load_workbook(str(path))

    if args.action == "info":
        sheets = []
        for name in wb.sheetnames:
            ws = wb[name]
            sheets.append(
                {
                    "name": name,
                    "max_row": ws.max_row,
                    "max_column": ws.max_column,
                }
            )
        print(json.dumps({"ok": True, "path": str(path.resolve()), "sheets": sheets}, ensure_ascii=False, indent=2))
        return 0

    sheet_name = args.sheet or wb.sheetnames[0]
    if sheet_name not in wb.sheetnames:
        print(json.dumps({"ok": False, "error": f"sheet not found: {sheet_name}"}))
        return 1
    ws = wb[sheet_name]

    if args.action == "read":
        rows = []
        for r in ws.iter_rows(
            min_row=1,
            max_row=min(ws.max_row or 1, args.max_rows),
            max_col=min(ws.max_column or 1, args.max_cols),
            values_only=True,
        ):
            rows.append([("" if v is None else v) for v in r])
        print(
            json.dumps(
                {"ok": True, "sheet": sheet_name, "rows": rows},
                ensure_ascii=False,
                indent=2,
                default=str,
            )
        )
        return 0

    if args.action == "write_cell":
        # try parse JSON value
        val: object = args.value
        try:
            val = json.loads(args.value)
        except Exception:
            pass
        ws[args.cell] = val
        out = Path(args.out) if args.out.strip() else path
        out.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(out))
        print(
            json.dumps(
                {"ok": True, "path": str(out.resolve()), "sheet": sheet_name, "cell": args.cell, "value": val},
                ensure_ascii=False,
                default=str,
            )
        )
        return 0

    print(json.dumps({"ok": False, "error": "unknown action"}))
    return 1


if __name__ == "__main__":
    sys.exit(main())
