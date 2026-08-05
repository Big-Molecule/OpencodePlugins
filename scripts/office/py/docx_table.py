#!/usr/bin/env python3
"""Add table or set cell text in a DOCX."""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path


def main() -> int:
    p = argparse.ArgumentParser()
    p.add_argument("docx")
    p.add_argument("--action", choices=["add", "set_cell", "to_csv"], required=True)
    p.add_argument("--rows", type=int, default=2)
    p.add_argument("--cols", type=int, default=2)
    p.add_argument("--data", default="", help='JSON 2D array for add, e.g. [["a","b"],["c","d"]]')
    p.add_argument("--table-index", type=int, default=0)
    p.add_argument("--row", type=int, default=0)
    p.add_argument("--col", type=int, default=0)
    p.add_argument("--text", default="")
    p.add_argument("--csv-out", default="")
    p.add_argument("--out", default="")
    args = p.parse_args()

    path = Path(args.docx)
    if not path.is_file():
        print(json.dumps({"ok": False, "error": f"not found: {path}"}))
        return 1

    from docx import Document

    doc = Document(str(path))
    out = Path(args.out) if args.out.strip() else path

    if args.action == "add":
        data = []
        if args.data.strip():
            data = json.loads(args.data)
        rows = max(args.rows, len(data) if data else 0, 1)
        cols = max(args.cols, max((len(r) for r in data), default=0), 1)
        table = doc.add_table(rows=rows, cols=cols)
        table.style = "Table Grid"
        for ri in range(min(rows, len(data))):
            for ci in range(min(cols, len(data[ri]))):
                table.rows[ri].cells[ci].text = str(data[ri][ci])
        out.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(out))
        print(json.dumps({"ok": True, "path": str(out.resolve()), "rows": rows, "cols": cols}, ensure_ascii=False))
        return 0

    if args.action == "set_cell":
        if args.table_index < 0 or args.table_index >= len(doc.tables):
            print(json.dumps({"ok": False, "error": f"table_index out of range 0..{len(doc.tables)-1}"}))
            return 1
        table = doc.tables[args.table_index]
        if args.row < 0 or args.row >= len(table.rows):
            print(json.dumps({"ok": False, "error": "row out of range"}))
            return 1
        if args.col < 0 or args.col >= len(table.rows[args.row].cells):
            print(json.dumps({"ok": False, "error": "col out of range"}))
            return 1
        table.rows[args.row].cells[args.col].text = args.text
        out.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(out))
        print(json.dumps({"ok": True, "path": str(out.resolve())}, ensure_ascii=False))
        return 0

    if args.action == "to_csv":
        if args.table_index < 0 or args.table_index >= len(doc.tables):
            print(json.dumps({"ok": False, "error": "table_index out of range"}))
            return 1
        import csv
        from io import StringIO

        table = doc.tables[args.table_index]
        buf = StringIO()
        w = csv.writer(buf)
        for row in table.rows:
            w.writerow([c.text.replace("\n", " ").strip() for c in row.cells])
        csv_text = buf.getvalue()
        if args.csv_out.strip():
            cp = Path(args.csv_out)
            cp.parent.mkdir(parents=True, exist_ok=True)
            cp.write_text(csv_text, encoding="utf-8")
            print(json.dumps({"ok": True, "csv_path": str(cp.resolve()), "rows": len(table.rows)}, ensure_ascii=False))
        else:
            print(json.dumps({"ok": True, "csv": csv_text, "rows": len(table.rows)}, ensure_ascii=False))
        return 0

    print(json.dumps({"ok": False, "error": "unknown action"}))
    return 1


if __name__ == "__main__":
    sys.exit(main())
