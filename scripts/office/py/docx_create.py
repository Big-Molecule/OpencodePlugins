#!/usr/bin/env python3
"""Create a simple DOCX from title + paragraphs JSON."""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path


def main() -> int:
    p = argparse.ArgumentParser()
    p.add_argument("out_docx")
    p.add_argument("--title", default="")
    p.add_argument(
        "--paragraphs",
        default="[]",
        help='JSON array of strings, e.g. ["a","b"]',
    )
    args = p.parse_args()

    try:
        paras = json.loads(args.paragraphs)
        if not isinstance(paras, list) or not all(isinstance(x, str) for x in paras):
            raise ValueError("paragraphs must be a JSON array of strings")
    except Exception as e:
        print(json.dumps({"ok": False, "error": f"bad --paragraphs: {e}"}))
        return 1

    from docx import Document
    from docx.shared import Pt

    doc = Document()
    if args.title.strip():
        h = doc.add_heading(args.title.strip(), level=1)
        for run in h.runs:
            run.font.size = Pt(16)
    for text in paras:
        if text.strip():
            doc.add_paragraph(text.strip())

    out = Path(args.out_docx)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    print(json.dumps({"ok": True, "path": str(out.resolve())}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
