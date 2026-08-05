#!/usr/bin/env python3
"""Set text of a paragraph by index (python-docx)."""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path


def main() -> int:
    p = argparse.ArgumentParser()
    p.add_argument("docx")
    p.add_argument("--index", type=int, required=True)
    p.add_argument("--text", required=True)
    p.add_argument("--out", default="", help="Output path (default: overwrite input)")
    args = p.parse_args()

    path = Path(args.docx)
    if not path.is_file():
        print(json.dumps({"ok": False, "error": f"not found: {path}"}))
        return 1

    from docx import Document

    doc = Document(str(path))
    if args.index < 0 or args.index >= len(doc.paragraphs):
        print(
            json.dumps(
                {
                    "ok": False,
                    "error": f"index {args.index} out of range 0..{len(doc.paragraphs)-1}",
                }
            )
        )
        return 1

    para = doc.paragraphs[args.index]
    # replace all runs with single run preserving style on paragraph
    if para.runs:
        para.runs[0].text = args.text
        for r in para.runs[1:]:
            r.text = ""
    else:
        para.add_run(args.text)

    out = Path(args.out) if args.out.strip() else path
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    print(
        json.dumps(
            {
                "ok": True,
                "path": str(out.resolve()),
                "index": args.index,
                "text_preview": args.text[:120],
            },
            ensure_ascii=False,
        )
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
