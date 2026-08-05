#!/usr/bin/env python3
"""Append body content of other DOCXs onto the first (simple merge)."""
from __future__ import annotations

import argparse
import json
import sys
from copy import deepcopy
from pathlib import Path


def main() -> int:
    p = argparse.ArgumentParser()
    p.add_argument("base_docx")
    p.add_argument("--append", action="append", default=[], help="DOCX paths to append (repeatable)")
    p.add_argument("--out", required=True)
    p.add_argument("--page-break", action="store_true", help="Insert page break before each append")
    args = p.parse_args()

    base = Path(args.base_docx)
    if not base.is_file():
        print(json.dumps({"ok": False, "error": f"not found: {base}"}))
        return 1
    if not args.append:
        print(json.dumps({"ok": False, "error": "provide at least one --append"}))
        return 1

    from docx import Document
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    def add_page_break(doc: Document) -> None:
        p = doc.add_paragraph()
        run = p.add_run()
        br = OxmlElement("w:br")
        br.set(qn("w:type"), "page")
        run._r.append(br)

    def append_doc(dst: Document, src_path: Path) -> None:
        src = Document(str(src_path))
        for element in src.element.body:
            # skip sectPr at end
            if element.tag == qn("w:sectPr"):
                continue
            dst.element.body.insert(-1, deepcopy(element))

    doc = Document(str(base))
    for ap in args.append:
        sp = Path(ap)
        if not sp.is_file():
            print(json.dumps({"ok": False, "error": f"append missing: {sp}"}))
            return 1
        if args.page_break:
            add_page_break(doc)
        append_doc(doc, sp)

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    print(
        json.dumps(
            {"ok": True, "path": str(out.resolve()), "appended": len(args.append)},
            ensure_ascii=False,
        )
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
